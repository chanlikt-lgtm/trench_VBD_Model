#!/usr/bin/env python3
"""Build a paper-ready Excel workbook from the P2.3-S0 certified JSON records.

One table per sheet (Q1..Q6) plus a notes sheet. Values are transcribed from the
certified artefacts; full-precision interval endpoints remain in the JSON files.
"""
from __future__ import annotations
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

D = Path(__file__).resolve().parent.parent / "p2_3_spatial_S0"
L = D / "levels"

A = {t: json.load(open(D / f"A_S0-{t}.json")) for t in ("L0", "L1", "L2")}
d01 = json.load(open(D / "DA_S0_01.json"))
d12 = json.load(open(D / "DA_S0_12.json"))
first = d12["dA01"] if "dA01" in d12 else d01["delta"]
S = {"L0": json.load(open(D / "CP-M2-C_L0.json")),
     "L1": json.load(open(D / "CP-L1_L1.json")),
     "L2": json.load(open(D / "CP-L2_L2.json"))}
M = {t: json.load(open(L / f"S0_{t}_mesh.json")) for t in ("L0", "L1", "L2")}
R = {"L0": json.load(open(D / "CP-M2-B_L0.json")),
     "L1": json.load(open(D / "CP-M2-B_L1.json")),
     "L2": json.load(open(D / "CP-L2_resources.json"))}
A6 = json.load(open(D / "CP-L2-A6_summary.json"))
MAN = json.load(open(D / "CP-M2-A_manifest.json"))

HDR = PatternFill("solid", fgColor="1F4E79")
HF = Font(bold=True, color="FFFFFF", size=11)
TITLE = Font(bold=True, size=13, color="1F4E79")
THIN = Border(*[Side(style="thin", color="B0B7BE")] * 4)
E15 = "0.000000000000000"
E9 = "0.000000000E+00"


def sheet(wb, name, title, headers, rows, widths=None, fmt=None):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = TITLE
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.fill = HDR; c.font = HF; c.border = THIN
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, row in enumerate(rows, start=4):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = THIN
            if fmt and j in fmt:
                c.number_format = fmt[j]
            c.alignment = Alignment(horizontal="center" if j > 1 else "left")
    for j, w in enumerate(widths or [28] * len(headers), start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    rows = [[f"S0-{t}", A[t]["lower"], A[t]["upper"], A[t]["width"],
             A[t]["charged"], A[t]["mixed"], A[t]["seconds"],
             A[t]["max_terminating_depth"],
             "BINDING" if A[t]["depth_limit_binding"] else "non-binding",
             A[t]["commit"]] for t in ("L0", "L1", "L2")]
    sheet(wb, "Q1_Areas",
          "Table Q1 — Certified neutral-area enclosures (AREA_TOL = 1e-13, um^2)",
          ["Level", "Lower bound", "Upper bound", "Width", "Charged elems", "Mixed elems",
           "Runtime (s)", "Max terminating depth", "Depth limit", "Commit"],
          rows, [10, 20, 20, 14, 13, 12, 12, 17, 13, 11],
          {2: E15, 3: E15, 4: E9})

    rows = [
        ["S0-L0 -> S0-L1", first[0], first[1],
         first[1] - first[0], "yes (positive)",
         f"{100*first[0]/A['L0']['upper']:.4f} – {100*first[1]/A['L0']['lower']:.4f} %"],
        ["S0-L1 -> S0-L2", d12["dA12"][0], d12["dA12"][1],
         d12["dA12"][1] - d12["dA12"][0], "yes (positive)",
         f"{100*d12['dA12'][0]/A['L1']['upper']:.4f} – {100*d12['dA12'][1]/A['L1']['lower']:.4f} %"],
    ]
    ws = sheet(wb, "Q2_Increments",
               "Table Q2 — Certified successive area increments (um^2)",
               ["Comparison", "Lower", "Upper", "Width", "Sign resolved",
                "Relative to coarser level"],
               rows, [18, 20, 20, 15, 16, 26], {2: E9, 3: E9, 4: E9})
    extra = [("Shrinkage test", ""),
             ("sup|dA_12|", d12["sup_abs_dA12"]),
             ("inf|dA_01|", d12["inf_abs_dA01"]),
             ("certified smaller", "TRUE" if d12["shrinkage"] else "FALSE"),
             ("ratio dA_12 / dA_01", f"{d12['ratio'][0]:.4f} – {d12['ratio'][1]:.4f}"),
             ("candidate p (descriptive only)", f"{d12['candidate_p'][0]:.3f} – {d12['candidate_p'][1]:.3f}"),
             ("caution", "two increments only; arc does not scale with r; nesting unverified")]
    for i, (k, v) in enumerate(extra, start=7):
        ws.cell(row=i, column=1, value=k).font = Font(bold=(k in ("Shrinkage test", "certified smaller")))
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = E9

    rows = []
    for t in ("L0", "L1", "L2"):
        s = S[t]; f = s["f7"]; o = s["osqp"]
        rows.append([f"S0-{t}", o["status"], o["iters"], o["obj"], o["pri"], o["dua"],
                     "identical" if s["checkpoint"]["identical"] else "MISMATCH",
                     s["active"], str(s.get("repairs")),
                     str(s.get("minres_info", "NOT_RECORDED")),
                     f["stationarity"], f["complementarity"], f["min_slack"],
                     f["violations"], f["lam_min"], f["oxide"], f["overshoot"],
                     "PASS" if f["gate"] else "FAIL"])
    sheet(wb, "Q3_F7", "Table Q3 — Solver statistics and full-set F7 certification",
          ["Level", "OSQP status", "Iterations", "Objective", "Primal res", "Dual res",
           "Checkpoint", "Active rows", "Repairs", "MINRES info", "Stationarity",
           "Complementarity", "Min slack", "Violations", "lam min", "Oxide",
           "Overshoot", "F7"],
          rows, [9, 12, 11, 16, 12, 12, 12, 11, 22, 13, 13, 15, 13, 11, 9, 12, 12, 8],
          {4: "0.000000E+00", 5: E9, 6: E9, 11: E9, 12: E9, 13: E9, 16: E9, 17: E9})

    rows = []
    for t in ("L0", "L1", "L2"):
        m = M[t]; r = R[t]; s = S[t]
        rows.append([f"S0-{t}", s["mesh_sha"], m["nodes"], m["elements"],
                     r["dofs_total"], r["dofs_free"], r["dofs_fixed"],
                     r["charged_elements"], r["bernstein_rows"],
                     m.get("min_angle"), m.get("max_aspect"),
                     r["seconds"]["total"], r["peak_rss_mb"],
                     r["K"]["nnz"], r["Kff"]["nnz"], r["C_F"]["nnz"]])
    sheet(wb, "Q4_Mesh",
          "Table Q4 — Mesh, degrees of freedom, constraints and assembly cost",
          ["Level", "Mesh SHA256", "Nodes", "Elements", "P2 DOFs (compacted)", "Free DOFs",
           "Fixed DOFs", "Charged elems", "Bernstein rows", "Min angle (deg)",
           "Max aspect", "Assembly (s)", "Assembly peak RSS (MB)", "K nnz", "Kff nnz", "C_F nnz"],
          rows, [9, 70, 11, 11, 17, 11, 11, 12, 15, 12, 11, 12, 18, 12, 12, 12],
          {12: "0.0", 13: "0"})

    rows = [[f"S0-{t}", A6["arc"][t][0], A6["arc"][t][1],
             A6["comp"][t][0], A6["comp"][t][1]] for t in ("L0", "L1", "L2")]
    ws = sheet(wb, "Q5_ArcRegion",
               "Table Q5 — Fixed arc-adjacent region diagnostic "
               "(R_arc: within 0.01 um of the canonical arcs)",
               ["Level", "Arc lower", "Arc upper", "Complement lower", "Complement upper"],
               rows, [10, 18, 18, 22, 22],
               {2: "0.000000E+00", 3: E9, 4: E15, 5: E15})
    ws.cell(row=8, column=1, value="Increments").font = Font(bold=True)
    for j, h in enumerate(["Quantity", "Lower", "Upper", "Contains zero"], start=1):
        ws.cell(row=9, column=j, value=h).font = Font(bold=True)
    data = [("dA_arc,01", A6["arc_d01"]), ("dA_arc,12", A6["arc_d12"]),
            ("dA_comp,01", A6["comp_d01"]), ("dA_comp,12", A6["comp_d12"]),
            ("dA_whole,01", A6["whole_d01"]), ("dA_whole,12", A6["whole_d12"])]
    for i, (n, v) in enumerate(data, start=10):
        ws.cell(row=i, column=1, value=n)
        ws.cell(row=i, column=2, value=v[0]).number_format = E9
        ws.cell(row=i, column=3, value=v[1]).number_format = E9
        ws.cell(row=i, column=4, value="yes" if v[0] <= 0 <= v[1] else "no")
    ws.cell(row=17, column=1,
            value="DIAGNOSTIC ONLY: arc lower bound is 0 at every level "
                  "(no fully neutral sub-triangle in the band); upper bounds are ceilings.")

    rows = []
    for t in ("L0", "L1", "L2"):
        mk = S[t]["marks"]
        rows.append([f"S0-{t}", mk["start"]["rss"], mk["after_assembly"]["rss"],
                     mk["after_osqp"]["rss"], mk["after_polish"]["rss"],
                     S[t]["peak_rss_mb"], min(v["avail"] for v in mk.values()),
                     round(mk["after_polish"]["t"] / 60, 1)])
    sheet(wb, "Q6_Memory", "Table Q6 — Measured memory and wall-clock trace",
          ["Level", "RSS start (MB)", "RSS after assembly", "RSS after OSQP",
           "RSS after polish", "Peak RSS (MB)", "Min available RAM (MB)", "Total (min)"],
          rows, [9, 15, 19, 16, 17, 14, 20, 11],
          {2: "0", 3: "0", 4: "0", 5: "0", 6: "0", 7: "0", 8: "0.0"})

    ws = wb.create_sheet("Notes")
    notes = [
        "P2.3-S0 controlled spatial-convergence study — certified results",
        "",
        "Model: conforming straight-sided P2 triangles; sharp-depletion obstacle via Bernstein",
        "positivity constraints; exact clipped partial-element donor integration at x = 0.54 um.",
        "Geometry: one frozen polygonal domain (G-fixed), four arcs at 64 chords each.",
        "Size field: h_k = r^k h0 with r = 2^(-1/2); global scaling of band, transition, tip",
        "and background. The fixed chord count is the declared geometric exception.",
        "",
        "Solver: checkpoint-first OSQP (eps_abs = eps_rel = 1e-8, max_iter 30000, scaling 10),",
        "then matrix-free scaled MINRES with residual correction and full-set F7 evaluation.",
        "OSQP 'solved' alone is not F7 certification.",
        "",
        "Areas: corrected Bernstein branch-and-bound enclosure at AREA_TOL = 1e-13,",
        "MAX_DEPTH = 24; the depth limit is non-binding at every level. Enclosures are",
        "one-sided bounds, not exact areas. Intervals are propagated by endpoint subtraction.",
        "",
        "Interpretation limits:",
        "  * three levels give two increments; no convergence order or continuum limit is claimed;",
        "  * the frozen polygon gives a polygonal-domain result, not the curved physical domain;",
        "  * L2 arc-adjacent elements refine at 0.784 of the prescribed ratio (target 0.7071);",
        "  * the direct local arc contribution is below 1.6e-7 um^2, but its nonlocal",
        "    (elliptic / active-set) effect is NOT bounded by this diagnostic;",
        "  * literal finite-element nesting of the levels has not been verified.",
        "",
        "Software:",
        f"  Python {MAN['software']['python']} · NumPy {MAN['software']['numpy']} · "
        f"SciPy {MAN['software']['scipy']} · OSQP {MAN['software']['osqp']} · "
        f"Gmsh {MAN['software']['gmsh']}",
        f"  {MAN['software']['platform']}",
        "",
        "Provenance commits:",
        "  e937926 design · 1c8d4df repaired pilots · a5b76e2 manifest",
        "  a3a5248 L0 F7 · 76f6f42 L0 area · e1c8a2d L1 F7 · 803dcff L1 area",
        "  c07ec2f L2 pilot · 6e8d9e1 L2 F7 · 5773158 L2 area/increments · ebff654 arc diagnostic",
    ]
    for i, line in enumerate(notes, start=1):
        c = ws.cell(row=i, column=1, value=line)
        if i == 1:
            c.font = TITLE
    ws.column_dimensions["A"].width = 108

    out = D / "P2_3_S0_Results_Tables.xlsx"
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
