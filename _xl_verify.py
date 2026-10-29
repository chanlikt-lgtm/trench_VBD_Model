from openpyxl import load_workbook
wb = load_workbook("p2_3_spatial_S0/P2_3_S0_Results_Tables.xlsx")
print("sheets:", wb.sheetnames)
for name in ("Q1_Areas","Q2_Increments","Q5_ArcRegion"):
    ws = wb[name]
    print(f"\n--- {name} ({ws.max_row}x{ws.max_column}) ---")
    for r in range(1, min(ws.max_row, 12)+1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column,7)+1)]
        if any(v is not None for v in vals):
            print(r, vals)
